"""Tests for arxiv_hunt.excel_converter module."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from arxiv_hunt.config import CSV_COLUMNS
from arxiv_hunt.csv_io import save_papers_to_csv
from arxiv_hunt.excel_converter import convert_csv_to_excel
from arxiv_hunt.models import Paper


def _prepare_csv(papers: list[Paper], tmp_path: Path) -> Path:
    """Helper: save papers to CSV and return the results.csv path."""
    csv_dir = tmp_path / "csv"
    csv_dir.mkdir(parents=True, exist_ok=True)
    results_path, _ = save_papers_to_csv(papers, "test", output_dir=csv_dir)
    return results_path


class TestConvertCsvToExcel:
    """Tests for convert_csv_to_excel."""

    def test_convert_creates_file(
        self, sample_papers: list[Paper], tmp_path: Path
    ) -> None:
        """An Excel file should be created at the specified output path."""
        csv_path = _prepare_csv(sample_papers, tmp_path)
        xlsx_path = tmp_path / "output.xlsx"

        result = convert_csv_to_excel(csv_path=csv_path, output_path=xlsx_path)

        assert result == xlsx_path
        assert xlsx_path.exists()

    def test_convert_header_row(
        self, sample_papers: list[Paper], tmp_path: Path
    ) -> None:
        """Header row should contain all CSV columns plus extra Excel columns."""
        csv_path = _prepare_csv(sample_papers, tmp_path)
        xlsx_path = tmp_path / "output.xlsx"
        convert_csv_to_excel(csv_path=csv_path, output_path=xlsx_path)

        wb = load_workbook(xlsx_path)
        ws = wb.active

        expected_columns = CSV_COLUMNS + ["abstract_translated", "relevance_score"]
        header_values = [ws.cell(row=1, column=i).value for i in range(1, len(expected_columns) + 1)]
        assert header_values == expected_columns

    def test_convert_data_rows(
        self, sample_papers: list[Paper], tmp_path: Path
    ) -> None:
        """Data rows should contain the correct paper data."""
        csv_path = _prepare_csv(sample_papers, tmp_path)
        xlsx_path = tmp_path / "output.xlsx"
        convert_csv_to_excel(csv_path=csv_path, output_path=xlsx_path)

        wb = load_workbook(xlsx_path)
        ws = wb.active

        # Row 2 = first data row; column 1 = arxiv_id
        assert ws.cell(row=2, column=1).value == sample_papers[0].arxiv_id
        # Row 3 = second data row
        assert ws.cell(row=3, column=1).value == sample_papers[1].arxiv_id

        # Verify title column (column 2)
        assert ws.cell(row=2, column=2).value == sample_papers[0].title

    def test_translate_formula(
        self, sample_papers: list[Paper], tmp_path: Path
    ) -> None:
        """The abstract_translated column should contain a TRANSLATE formula."""
        csv_path = _prepare_csv(sample_papers, tmp_path)
        xlsx_path = tmp_path / "output.xlsx"
        convert_csv_to_excel(csv_path=csv_path, output_path=xlsx_path)

        wb = load_workbook(xlsx_path)
        ws = wb.active

        # abstract_translated column index (1-based)
        all_columns = CSV_COLUMNS + ["abstract_translated", "relevance_score"]
        translated_col = all_columns.index("abstract_translated") + 1

        cell_value = ws.cell(row=2, column=translated_col).value
        assert cell_value is not None
        assert "TRANSLATE" in cell_value
        assert '"en"' in cell_value
        assert '"ja"' in cell_value

    def test_frozen_panes(
        self, sample_papers: list[Paper], tmp_path: Path
    ) -> None:
        """Panes should be frozen at A2 (top row frozen)."""
        csv_path = _prepare_csv(sample_papers, tmp_path)
        xlsx_path = tmp_path / "output.xlsx"
        convert_csv_to_excel(csv_path=csv_path, output_path=xlsx_path)

        wb = load_workbook(xlsx_path)
        ws = wb.active

        assert ws.freeze_panes == "A2"
