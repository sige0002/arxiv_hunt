"""Excel converter with TRANSLATE formula embedding."""

from __future__ import annotations

import logging
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from arxiv_hunt.config import CSV_COLUMNS, CSV_OUTPUT_DIR, EXCEL_OUTPUT_DIR, RESULTS_CSV_FILENAME
from arxiv_hunt.csv_io import load_papers_from_csv

logger = logging.getLogger(__name__)

# Header styling
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_HEADER_FILL = PatternFill(start_color="B31B1B", end_color="B31B1B", fill_type="solid")
_HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)

# Column widths
_COLUMN_WIDTHS: dict[str, int] = {
    "arxiv_id": 18,
    "title": 50,
    "authors": 30,
    "abstract": 60,
    "published": 12,
    "updated": 12,
    "primary_category": 15,
    "categories": 25,
    "pdf_url": 30,
    "entry_url": 30,
    "abstract_translated": 60,
    "relevance_score": 15,
}


def convert_csv_to_excel(
    csv_path: Path | None = None,
    output_path: Path | None = None,
    *,
    source_lang: str = "en",
    target_lang: str = "ja",
) -> Path:
    """Convert a CSV file to Excel with TRANSLATE formulas.

    Args:
        csv_path: Input CSV path (defaults to data/csv/results.csv).
        output_path: Output Excel path (defaults to data/excel/results.xlsx).
        source_lang: Source language code for TRANSLATE formula.
        target_lang: Target language code for TRANSLATE formula.

    Returns:
        Path to the generated Excel file.
    """
    csv_path = csv_path or (CSV_OUTPUT_DIR / RESULTS_CSV_FILENAME)
    if output_path is None:
        EXCEL_OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
        output_path = EXCEL_OUTPUT_DIR / csv_path.with_suffix(".xlsx").name

    papers = load_papers_from_csv(csv_path)

    wb = Workbook()
    ws = wb.active
    assert ws is not None  # Workbook always has an active sheet
    ws.title = "Papers"

    # All columns: CSV columns + abstract_translated + relevance_score
    all_columns = CSV_COLUMNS + ["abstract_translated", "relevance_score"]

    # Write headers
    for col_idx, col_name in enumerate(all_columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGNMENT

    # Find abstract column index (1-based)
    abstract_col_idx = CSV_COLUMNS.index("abstract") + 1
    abstract_translated_col_idx = all_columns.index("abstract_translated") + 1

    # Write data rows
    for row_idx, paper in enumerate(papers, start=2):
        row_data = paper.to_csv_row()
        for col_idx, col_name in enumerate(CSV_COLUMNS, start=1):
            ws.cell(row=row_idx, column=col_idx, value=row_data[col_name])

        # TRANSLATE formula for abstract
        abstract_cell_ref = f"{get_column_letter(abstract_col_idx)}{row_idx}"
        formula = f'=_xlfn.TRANSLATE({abstract_cell_ref},"{source_lang}","{target_lang}")'
        ws.cell(row=row_idx, column=abstract_translated_col_idx, value=formula)

        # relevance_score left empty

    # Set column widths
    for col_idx, col_name in enumerate(all_columns, start=1):
        width = _COLUMN_WIDTHS.get(col_name, 15)
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Freeze top row
    ws.freeze_panes = "A2"

    # Auto-filter
    last_col_letter = get_column_letter(len(all_columns))
    last_row = max(ws.max_row, 2)
    ws.auto_filter.ref = f"A1:{last_col_letter}{last_row}"

    # Wrap text for abstract columns
    for row_idx in range(2, len(papers) + 2):
        ws.cell(row=row_idx, column=abstract_col_idx).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=row_idx, column=abstract_translated_col_idx).alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(output_path)
    logger.info("Saved Excel to %s (%d papers)", output_path, len(papers))

    return output_path
